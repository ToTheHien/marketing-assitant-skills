"""
Tạo dummy template files (.docx) cho skill investment-progress-summary.
Chạy: python create_dummy_templates.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(SCRIPT_DIR, "..", "templates")
os.makedirs(TEMPLATES_DIR, exist_ok=True)


def set_cell_background(cell, fill_color: str):
    """Set cell background color (hex without #)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_color)
    tcPr.append(shd)


def style_header_row(row, bg_color="2E75B6"):
    """Style a table row as header: white bold text, colored background."""
    for cell in row.cells:
        set_cell_background(cell, bg_color)
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def add_placeholder_para(doc, text: str, italic=True, color="808080"):
    p = doc.add_paragraph()
    run = p.add_run(f"[{text}]")
    run.italic = italic
    r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
    run.font.color.rgb = RGBColor(r, g, b)
    return p


def create_weekly_template():
    doc = Document()

    # Page margins (A4, 2cm all sides)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    # --- HEADER ---
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header_para.add_run("BECAMEX IDC — KHU CÔNG NGHIỆP")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("BÁO CÁO TIẾN ĐỘ ĐẦU TƯ TUẦN")
    run.bold = True
    run.font.size = Pt(16)

    meta_para = doc.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_para.add_run("Kỳ báo cáo: ")
    run = meta_para.add_run("[Tuần XX, DD/MM – DD/MM/YYYY]")
    run.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- WATERMARK note ---
    watermark = doc.add_paragraph()
    watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = watermark.add_run("⚠️  TEMPLATE MẪU – CHƯA CÓ DỮ LIỆU THỰC TẾ  ⚠️")
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    doc.add_paragraph()

    # --- SECTION 1: Executive Summary ---
    h1 = doc.add_heading("I. TÓM TẮT ĐIỀU HÀNH", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    bullets = [
        "Điểm nổi bật 1: [Mô tả milestone hoặc sự kiện quan trọng]",
        "Điểm nổi bật 2: [Cập nhật tình hình giải ngân]",
        "Điểm nổi bật 3: [Vướng mắc đã xử lý hoặc đang xử lý]",
    ]
    for b in bullets:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(b)
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- SECTION 2: Overview Table ---
    h2 = doc.add_heading("II. TỔNG QUAN TIẾN ĐỘ", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    headers = ["STT", "Nhà đầu tư", "Dự án / Lô đất", "Tiến độ (%)", "Giải ngân (tỷ đ)", "Trạng thái"]
    widths = [0.6, 2.5, 3.0, 1.2, 1.8, 1.5]
    table = doc.add_table(rows=3, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, (h, w) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[i]
        cell.width = Inches(w)
        cell.paragraphs[0].add_run(h).bold = True
    style_header_row(table.rows[0])

    # Dummy rows
    dummy_data = [
        ["1", "[Tên NĐT A]", "[Tên dự án / Lô XX]", "35%", "12.5", "Đúng tiến độ"],
        ["2", "[Tên NĐT B]", "[Tên dự án / Lô YY]", "20%", "8.0", "Chậm nhẹ"],
    ]
    for row_idx, row_data in enumerate(dummy_data):
        row = table.rows[row_idx + 1]
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            run = cell.paragraphs[0].add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- SECTION 3: Detail per investor ---
    h3 = doc.add_heading("III. CẬP NHẬT CHI TIẾT THEO NHÀ ĐẦU TƯ", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    for inv in ["A", "B"]:
        sub = doc.add_heading(f"3.{inv}. [Tên nhà đầu tư {inv}]", level=2)
        sub.runs[0].font.color.rgb = RGBColor(0x17, 0x50, 0x8D)

        info_items = [
            ("Dự án", "[Tên dự án]"),
            ("Vốn đăng ký", "[XX triệu USD]"),
            ("Tiến độ chung", "[XX%]"),
        ]
        for label, val in info_items:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            run = p.add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph("Tuần này:").runs[0].bold = True
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run("[Công việc đã thực hiện / milestone đạt được]")
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        p = doc.add_paragraph()
        p.add_run("Vướng mắc: ").bold = True
        run = p.add_run("[Không có / Mô tả vướng mắc nếu có]")
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph()

    # --- SECTION 4: Issues ---
    h4 = doc.add_heading("IV. VƯỚNG MẮC & KIẾN NGHỊ", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    issue_headers = ["Nhà đầu tư", "Vướng mắc", "Người phụ trách", "Deadline"]
    issue_table = doc.add_table(rows=2, cols=4)
    issue_table.style = "Table Grid"
    for i, h in enumerate(issue_headers):
        cell = issue_table.rows[0].cells[i]
        cell.paragraphs[0].add_run(h).bold = True
    style_header_row(issue_table.rows[0])

    dummy_issue = ["[NĐT A]", "[Mô tả vướng mắc]", "[Tên người phụ trách]", "[DD/MM/YYYY]"]
    for i, val in enumerate(dummy_issue):
        cell = issue_table.rows[1].cells[i]
        run = cell.paragraphs[0].add_run(val)
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- SECTION 5: Next Week Plan ---
    h5 = doc.add_heading("V. KẾ HOẠCH TUẦN TỚI", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    plan_items = [
        "[NĐT A]: [Công việc dự kiến tuần tới]",
        "[NĐT B]: [Công việc dự kiến tuần tới]",
        "[Các cuộc họp/kiểm tra đã lên lịch]",
    ]
    for item in plan_items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- FOOTER ---
    doc.add_paragraph("─" * 60)
    footer_para = doc.add_paragraph()
    footer_para.add_run("Người lập báo cáo: ").bold = True
    footer_para.add_run("_______________________")
    footer_para.add_run("          Ngày: ").bold = True
    footer_para.add_run("DD/MM/YYYY")

    out_path = os.path.join(TEMPLATES_DIR, "weekly_template.docx")
    doc.save(out_path)
    print(f"Created: {out_path}")
    return out_path


def create_monthly_template():
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2)

    # --- COVER PAGE ---
    for _ in range(5):
        doc.add_paragraph()

    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_title.add_run("BÁO CÁO TIẾN ĐỘ ĐẦU TƯ")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    cover_sub = doc.add_paragraph()
    cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_sub.add_run("BECAMEX IDC – KHU CÔNG NGHIỆP")
    run.font.size = Pt(14)
    run.bold = True

    cover_month = doc.add_paragraph()
    cover_month.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_month.add_run("[Tháng MM/YYYY]")
    run.font.size = Pt(14)
    run.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for _ in range(3):
        doc.add_paragraph()

    watermark = doc.add_paragraph()
    watermark.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = watermark.add_run("⚠️  TEMPLATE MẪU – CHƯA CÓ DỮ LIỆU THỰC TẾ  ⚠️")
    run.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    cover_unit = doc.add_paragraph()
    cover_unit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_unit.add_run("Đơn vị lập: [Tên phòng/ban]")
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for _ in range(3):
        doc.add_paragraph()

    confidential = doc.add_paragraph()
    confidential.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = confidential.add_run("PHÂN LOẠI: NỘI BỘ")
    run.bold = True
    run.font.size = Pt(11)

    doc.add_page_break()

    # --- SECTION 1: Executive Summary ---
    h1 = doc.add_heading("I. TÓM TẮT ĐIỀU HÀNH", level=1)
    h1.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    kpi_items = [
        ("Tổng số nhà đầu tư đang theo dõi", "[X dự án]"),
        ("Tổng vốn đăng ký", "[XXX triệu USD]"),
        ("Tổng giải ngân lũy kế", "[XXX triệu USD (XX%)]"),
        ("Giải ngân trong tháng", "[XX triệu USD]"),
    ]
    for label, val in kpi_items:
        p = doc.add_paragraph()
        p.add_run(f"• {label}: ").bold = True
        run = p.add_run(val)
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph("Điểm nổi bật tháng:").runs[0].bold = True
    for item in ["[Điểm nổi bật 1]", "[Điểm nổi bật 2]", "[Điểm nổi bật 3]"]:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- SECTION 2: Portfolio Overview ---
    h2 = doc.add_heading("II. TỔNG QUAN DANH MỤC ĐẦU TƯ", level=1)
    h2.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    headers = ["STT", "Nhà đầu tư", "Quốc gia", "Ngành", "Vốn đăng ký\n(tr.USD)", "G.ngân lũy kế (%)", "Tiến độ XD (%)", "Trạng thái"]
    table = doc.add_table(rows=4, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].add_run(h).bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
    style_header_row(table.rows[0])

    dummy_rows = [
        ["1", "[NĐT A]", "[Hàn Quốc]", "[Điện tử]", "50", "60%", "45%", "Đúng tiến độ"],
        ["2", "[NĐT B]", "[Nhật Bản]", "[Cơ khí]", "30", "40%", "30%", "Chậm nhẹ"],
        ["3", "[NĐT C]", "[Việt Nam]", "[Logistics]", "20", "80%", "75%", "Đúng tiến độ"],
    ]
    for row_idx, row_data in enumerate(dummy_rows):
        row = table.rows[row_idx + 1]
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            run = cell.paragraphs[0].add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            run.font.size = Pt(9)

    doc.add_paragraph()

    # --- SECTION 3: Detail per investor ---
    h3 = doc.add_heading("III. TIẾN ĐỘ CHI TIẾT TỪNG NHÀ ĐẦU TƯ", level=1)
    h3.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    for inv_label in ["A", "B", "C"]:
        sub = doc.add_heading(f"3.{inv_label}. [Tên nhà đầu tư {inv_label}]", level=2)
        sub.runs[0].font.color.rgb = RGBColor(0x17, 0x50, 0x8D)

        fields = [
            ("Dự án", "[Tên dự án]"),
            ("Diện tích", "[XX ha]"),
            ("Vốn đăng ký", "[XX triệu USD]"),
            ("Giải ngân lũy kế", "[XX triệu USD (XX%)]"),
            ("Tiến độ xây dựng", "[XX%]"),
        ]
        for label, val in fields:
            p = doc.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            run = p.add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph("Milestone tháng này:").runs[0].bold = True
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run("[Milestone đã đạt được]")
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        p = doc.add_paragraph()
        p.add_run("Vướng mắc: ").bold = True
        run = p.add_run("[Không có / Mô tả vướng mắc nếu có]")
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph()

    # --- SECTION 4: Legal/Permits ---
    h4 = doc.add_heading("IV. TÌNH HÌNH PHÁP LÝ / THỦ TỤC", level=1)
    h4.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    legal_headers = ["Nhà đầu tư", "Loại hồ sơ", "Cơ quan cấp", "Trạng thái", "Ghi chú"]
    legal_table = doc.add_table(rows=3, cols=5)
    legal_table.style = "Table Grid"
    for i, h in enumerate(legal_headers):
        cell = legal_table.rows[0].cells[i]
        cell.paragraphs[0].add_run(h).bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
    style_header_row(legal_table.rows[0])

    for r in range(1, 3):
        for c, val in enumerate(["[NĐT]", "[Giấy phép XD / PCCC / ...]", "[Sở Xây dựng / ...]", "[Đang xử lý / Đã cấp]", "[Ghi chú]"]):
            run = legal_table.rows[r].cells[c].paragraphs[0].add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            run.font.size = Pt(9)

    doc.add_paragraph()

    # --- SECTION 5: Financial ---
    h5 = doc.add_heading("V. TÌNH HÌNH TÀI CHÍNH", level=1)
    h5.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    fin_headers = ["Nhà đầu tư", "Vốn cam kết\n(tr.USD)", "G.ngân lũy kế\n(tr.USD)", "Tháng này\n(tr.USD)", "Dự kiến tháng tới\n(tr.USD)"]
    fin_table = doc.add_table(rows=4, cols=5)
    fin_table.style = "Table Grid"
    for i, h in enumerate(fin_headers):
        cell = fin_table.rows[0].cells[i]
        cell.paragraphs[0].add_run(h).bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
    style_header_row(fin_table.rows[0])

    for r in range(1, 3):
        for c, val in enumerate(["[NĐT]", "[XX]", "[XX]", "[XX]", "[XX]"]):
            run = fin_table.rows[r].cells[c].paragraphs[0].add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            run.font.size = Pt(9)

    # Total row
    for c, val in enumerate(["TỔNG", "[XXX]", "[XXX]", "[XX]", "[XX]"]):
        cell = fin_table.rows[3].cells[c]
        run = cell.paragraphs[0].add_run(val)
        run.bold = True
        run.font.size = Pt(9)
    style_header_row(fin_table.rows[3], bg_color="D6E4F0")
    for cell in fin_table.rows[3].cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)

    doc.add_paragraph()

    # --- SECTION 6: Risk ---
    h6 = doc.add_heading("VI. RỦI RO VÀ BIỆN PHÁP", level=1)
    h6.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    risk_headers = ["Rủi ro", "Mức độ", "NĐT liên quan", "Biện pháp", "Người p.trách"]
    risk_table = doc.add_table(rows=3, cols=5)
    risk_table.style = "Table Grid"
    for i, h in enumerate(risk_headers):
        cell = risk_table.rows[0].cells[i]
        cell.paragraphs[0].add_run(h).bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
    style_header_row(risk_table.rows[0])

    for r in range(1, 3):
        for c, val in enumerate(["[Mô tả rủi ro]", "Trung bình / Cao", "[NĐT]", "[Biện pháp]", "[Tên]"]):
            run = risk_table.rows[r].cells[c].paragraphs[0].add_run(val)
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
            run.font.size = Pt(9)

    doc.add_paragraph()

    # --- SECTION 7: Next Month Plan ---
    h7 = doc.add_heading("VII. KẾ HOẠCH THÁNG TỚI", level=1)
    h7.runs[0].font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

    p = doc.add_paragraph()
    run = p.add_run("Mục tiêu tháng [MM+1/YYYY]:")
    run.bold = True

    for item in ["[NĐT A]: Mục tiêu / công việc dự kiến", "[NĐT B]: Mục tiêu / công việc dự kiến", "[Sự kiện / cuộc họp quan trọng]"]:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        run.italic = True
        run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()

    # --- FOOTER ---
    doc.add_paragraph("─" * 60)
    footer = doc.add_paragraph()
    footer.add_run("Người lập: ").bold = True
    footer.add_run("_______________________")
    footer.add_run("     Phê duyệt: ").bold = True
    footer.add_run("_______________________")
    footer.add_run("     Ngày: ").bold = True
    footer.add_run("DD/MM/YYYY")

    out_path = os.path.join(TEMPLATES_DIR, "monthly_template.docx")
    doc.save(out_path)
    print(f"Created: {out_path}")
    return out_path


if __name__ == "__main__":
    print("Creating dummy investment progress report templates...")
    w = create_weekly_template()
    m = create_monthly_template()
    print(f"\nDone! Templates created in: {os.path.abspath(TEMPLATES_DIR)}")
    print("  - weekly_template.docx")
    print("  - monthly_template.docx")
    print("\nNote: These are placeholder templates. Replace with real branded templates when available.")
