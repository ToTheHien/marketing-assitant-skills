"""
Tạo dummy template files (.docx) cho skill promote-activity-summary.
Chạy từ project root:
  python .claude/skills/promote-activity-summary/scripts/create_dummy_templates.py
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATES_DIR = os.path.join(SCRIPT_DIR, "..", "templates")
os.makedirs(TEMPLATES_DIR, exist_ok=True)

BECAMEX_BLUE = "1F4E79"
ACCENT_BLUE = "2E75B6"
LIGHT_BLUE = "D6E4F0"
LIGHT_GRAY = "F2F2F2"
RED = "FF0000"


def set_cell_bg(cell, color_hex: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def run_color(run, color_hex: str):
    run.font.color.rgb = RGBColor(
        int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
    )


def header_row(table, texts: list, bg=BECAMEX_BLUE):
    row = table.rows[0]
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = ""
        set_cell_bg(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10)
        run_color(run, "FFFFFF")


def add_data_row(table, texts: list, bg=None):
    row = table.add_row()
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = ""
        if bg:
            set_cell_bg(cell, bg)
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.font.size = Pt(10)
    return row


def section_heading(doc, text: str):
    p = doc.add_paragraph()
    p.style = "Heading 1"
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run_color(run, BECAMEX_BLUE)
    return p


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p


def set_col_widths(table, widths_cm: list):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i < len(widths_cm):
                cell.width = Cm(widths_cm[i])


# ─────────────────────────── WEEKLY TEMPLATE ─────────────────────────────
def create_weekly(path: str):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    # Title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BECAMEX IDC — KHU CÔNG NGHIỆP")
    run.bold = True
    run.font.size = Pt(14)
    run_color(run, BECAMEX_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BÁO CÁO HOẠT ĐỘNG XÚC TIẾN ĐẦU TƯ TUẦN")
    run.bold = True
    run.font.size = Pt(16)
    run_color(run, BECAMEX_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Kỳ báo cáo: ")
    run.font.size = Pt(11)
    run2 = p.add_run("[Tuần XX, DD/MM – DD/MM/YYYY]")
    run2.font.size = Pt(11)
    run2.bold = True
    run_color(run2, ACCENT_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("⚠️  TEMPLATE MẪU – CHƯA CÓ DỮ LIỆU THỰC TẾ  ⚠️")
    run.bold = True
    run.font.size = Pt(11)
    run_color(run, RED)

    doc.add_paragraph()

    # Section 1
    section_heading(doc, "I. TÓM TẮT ĐIỀU HÀNH")
    p = doc.add_paragraph()
    run = p.add_run("Số liệu nhanh trong tuần:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "Số cuộc gặp nhà đầu tư tiềm năng: [X]")
    add_bullet(doc, "Số hoạt động phòng thương mại/hiệp hội: [X]")
    add_bullet(doc, "Số sự kiện tham dự/tổ chức: [X]")
    p = doc.add_paragraph()
    run = p.add_run("Điểm nổi bật:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[Điểm nổi bật 1 – cuộc gặp quan trọng nhất]")
    add_bullet(doc, "[Điểm nổi bật 2 – kết quả phòng thương mại]")
    add_bullet(doc, "[Điểm nổi bật 3 – sự kiện nổi bật]")

    # Section 2
    section_heading(doc, "II. GẶP GỠ NHÀ ĐẦU TƯ TIỀM NĂNG")
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    header_row(t, ["STT", "Tên công ty", "Quốc gia", "Ngành", "Hình thức", "Trạng thái pipeline", "Người phụ trách"])
    set_col_widths(t, [1.0, 4.0, 2.0, 2.5, 2.5, 3.5, 3.0])
    add_data_row(t, ["1", "[Tên công ty A]", "[QG]", "[Ngành]", "Site visit", "Đang quan tâm", "[Tên]"], LIGHT_BLUE)
    add_data_row(t, ["2", "[Tên công ty B]", "[QG]", "[Ngành]", "Online meeting", "Mới tiếp cận", "[Tên]"])

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("Chi tiết cuộc gặp đáng chú ý:")
    run.bold = True
    run.font.size = Pt(11)
    run_color(run, ACCENT_BLUE)
    add_bullet(doc, "Công ty: [Tên] | Quy mô quan tâm: [XX ha / XX m²]")
    add_bullet(doc, "Nội dung: [Tóm tắt cuộc trao đổi]")
    add_bullet(doc, "Kết quả: [Đồng ý site visit / Yêu cầu thêm thông tin / ...]")
    add_bullet(doc, "Bước tiếp theo: [Gửi hồ sơ / Lên lịch site visit] — Deadline: DD/MM/YYYY")

    # Section 3
    section_heading(doc, "III. HOẠT ĐỘNG PHÒNG THƯƠNG MẠI & HIỆP HỘI")
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    header_row(t, ["Tổ chức", "Loại hoạt động", "Ngày", "Kết quả", "Follow-up"])
    set_col_widths(t, [3.0, 3.5, 2.0, 4.0, 4.0])
    add_data_row(t, ["[KCCI / KOCHAM / JETRO / EuroCham / ...]", "[Loại HĐ]", "DD/MM", "[Kết quả]", "[Follow-up]"], LIGHT_BLUE)

    # Section 4
    section_heading(doc, "IV. SỰ KIỆN XÚC TIẾN ĐẦU TƯ")
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    header_row(t, ["Tên sự kiện", "Địa điểm", "Vai trò Becamex", "Số contacts", "Leads tiềm năng"])
    set_col_widths(t, [4.0, 3.0, 3.0, 2.5, 4.0])
    add_data_row(t, ["[Tên sự kiện]", "[Địa điểm]", "Tổ chức / Tham dự", "[X]", "[Tên công ty...]"], LIGHT_BLUE)

    # Section 5
    section_heading(doc, "V. KẾ HOẠCH TUẦN TỚI")
    p = doc.add_paragraph()
    run = p.add_run("Các cuộc gặp đã lên lịch:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "DD/MM: [Tên công ty] — [Hình thức] — [Người phụ trách]")
    p = doc.add_paragraph()
    run = p.add_run("Sự kiện:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[Tên sự kiện] — DD/MM — [Địa điểm]")
    p = doc.add_paragraph()
    run = p.add_run("Công việc cần hoàn thành:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[Gửi proposal / Chuẩn bị tài liệu / ...]")

    # Footer line
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Người lập: _____________      Ngày lập: ___/___/______      Phân phối: Ban lãnh đạo")
    run.font.size = Pt(9)
    run_color(run, "595959")

    doc.save(path)
    print(f"Created: {path}")


# ─────────────────────────── MONTHLY TEMPLATE ────────────────────────────
def create_monthly(path: str):
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2)

    # Cover page
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BECAMEX IDC — KHU CÔNG NGHIỆP")
    run.bold = True
    run.font.size = Pt(16)
    run_color(run, BECAMEX_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BÁO CÁO HOẠT ĐỘNG XÚC TIẾN ĐẦU TƯ")
    run.bold = True
    run.font.size = Pt(20)
    run_color(run, BECAMEX_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Tháng [MM/YYYY]")
    run.bold = True
    run.font.size = Pt(16)
    run_color(run, ACCENT_BLUE)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("⚠️  TEMPLATE MẪU – CHƯA CÓ DỮ LIỆU THỰC TẾ  ⚠️")
    run.bold = True
    run.font.size = Pt(12)
    run_color(run, RED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Đơn vị lập: [Tên phòng/ban]  |  Phân loại: Nội bộ")
    run.font.size = Pt(11)
    run_color(run, "595959")

    doc.add_page_break()

    # Section 1 – Summary
    section_heading(doc, "I. TÓM TẮT THÁNG")
    p = doc.add_paragraph()
    run = p.add_run("Dashboard số liệu:")
    run.bold = True
    run.font.size = Pt(11)

    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    cells = [
        ("Tổng cuộc gặp NĐT tiềm năng", "[X cuộc gặp]"),
        ("NĐT tiềm năng mới tiếp cận", "[X công ty]"),
        ("Dự án đang đàm phán", "[X dự án]"),
        ("Sự kiện tổ chức/tham dự", "[X sự kiện]"),
    ]
    for r in range(2):
        for c in range(2):
            idx = r * 2 + c
            cell = t.rows[r].cells[c]
            cell.text = ""
            set_cell_bg(cell, LIGHT_BLUE)
            p2 = cell.paragraphs[0]
            lbl = p2.add_run(cells[idx][0] + "\n")
            lbl.bold = True
            lbl.font.size = Pt(10)
            val = p2.add_run(cells[idx][1])
            val.bold = True
            val.font.size = Pt(14)
            run_color(val, BECAMEX_BLUE)

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run("Điểm nổi bật tháng:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[Điểm nổi bật 1]")
    add_bullet(doc, "[Điểm nổi bật 2]")
    add_bullet(doc, "[Điểm nổi bật 3]")

    # Section 2 – Pipeline
    section_heading(doc, "II. PIPELINE NHÀ ĐẦU TƯ TIỀM NĂNG")
    t = doc.add_table(rows=1, cols=8)
    t.style = "Table Grid"
    header_row(t, ["STT", "Tên công ty", "QG", "Ngành", "Quy mô QT", "Trạng thái", "Người PT", "Ghi chú"])
    set_col_widths(t, [0.8, 3.5, 1.5, 2.0, 2.5, 2.5, 2.0, 2.5])
    add_data_row(t, ["1", "[Công ty A]", "[QG]", "[Ngành]", "[XX ha]", "Đang đàm phán", "[Tên]", ""], LIGHT_BLUE)
    add_data_row(t, ["2", "[Công ty B]", "[QG]", "[Ngành]", "[XX m²]", "Đang quan tâm", "[Tên]", ""])

    # Section 3 – Detail per investor
    section_heading(doc, "III. CHI TIẾT HOẠT ĐỘNG THEO NHÀ ĐẦU TƯ")
    p = doc.add_paragraph()
    run = p.add_run("[Công ty A] — [Quốc gia] — [Ngành]")
    run.bold = True
    run.font.size = Pt(12)
    run_color(run, ACCENT_BLUE)
    add_bullet(doc, "Lịch sử tiếp xúc tháng: [Tóm tắt các cuộc gặp]")
    add_bullet(doc, "Trạng thái: [Đang đàm phán]")
    add_bullet(doc, "Bước tiếp theo: [Hành động cụ thể] — Deadline: DD/MM/YYYY")

    # Section 4 – Chambers
    section_heading(doc, "IV. HOẠT ĐỘNG PHÒNG THƯƠNG MẠI & HIỆP HỘI")
    t = doc.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    header_row(t, ["Tổ chức", "Số hoạt động", "Nội dung chính", "Kết quả nổi bật"])
    set_col_widths(t, [3.0, 2.5, 5.0, 6.0])
    add_data_row(t, ["KCCI", "[X]", "[Nội dung]", "[Kết quả]"], LIGHT_BLUE)
    add_data_row(t, ["KOCHAM", "[X]", "[Nội dung]", "[Kết quả]"])
    add_data_row(t, ["JETRO", "[X]", "[Nội dung]", "[Kết quả]"], LIGHT_BLUE)

    # Section 5 – Events
    section_heading(doc, "V. SỰ KIỆN XÚC TIẾN ĐẦU TƯ")
    t = doc.add_table(rows=1, cols=7)
    t.style = "Table Grid"
    header_row(t, ["Tên sự kiện", "Ngày", "Địa điểm", "Vai trò", "Số tham dự", "Leads", "Đánh giá"])
    set_col_widths(t, [3.5, 1.5, 2.5, 2.0, 2.0, 2.0, 3.0])
    add_data_row(t, ["[Tên sự kiện]", "DD/MM", "[Địa điểm]", "Tổ chức", "[X]", "[X]", "[Đánh giá]"], LIGHT_BLUE)

    # Section 6 – Analysis
    section_heading(doc, "VI. PHÂN TÍCH & ĐÁNH GIÁ THÁNG")
    add_bullet(doc, "Quốc gia/ngành có nhiều quan tâm nhất: [Phân tích]")
    add_bullet(doc, "Kênh tiếp cận hiệu quả nhất: [Phân tích]")
    add_bullet(doc, "Thách thức: [Mô tả]")
    add_bullet(doc, "Đề xuất cải thiện: [Đề xuất]")

    # Section 7 – Plan
    section_heading(doc, "VII. KẾ HOẠCH THÁNG TỚI")
    p = doc.add_paragraph()
    run = p.add_run("Mục tiêu:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[X cuộc gặp, X sự kiện, X tổ chức tiếp cận]")
    p = doc.add_paragraph()
    run = p.add_run("Hoạt động đã lên lịch:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[Ngày] [Tên hoạt động] — [Người phụ trách]")
    p = doc.add_paragraph()
    run = p.add_run("Sự kiện lớn:")
    run.bold = True
    run.font.size = Pt(11)
    add_bullet(doc, "[Tên sự kiện] — [Ngày] — [Địa điểm]")

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Người lập: _____________   Ngày: ___/___/______   Phê duyệt: _____________")
    run.font.size = Pt(9)
    run_color(run, "595959")

    doc.save(path)
    print(f"Created: {path}")


if __name__ == "__main__":
    print("Creating dummy templates for promote-activity-summary...")
    create_weekly(os.path.join(TEMPLATES_DIR, "weekly_template.docx"))
    create_monthly(os.path.join(TEMPLATES_DIR, "monthly_template.docx"))
    print("Done. Templates saved to:", TEMPLATES_DIR)
